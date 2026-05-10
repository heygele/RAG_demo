import docx
import os


def load_document(path: str):
    ext = os.path.splitext(path)[1].lower()
    if ext == '.docx':
        document = docx.Document(path)
        return "\n".join(para.text for para in document.paragraphs)
    elif ext == '.doc':
        document = docx.Document(path)
    elif ext == ".pdf":
        pass

    pass


if __name__ == "__main__":
    path = r"C:\Users\HP\Desktop\毕业论文修改指南-葛乐-车辆检测与流量计数系统设计与实现.docx"
    text = load_document(path)
